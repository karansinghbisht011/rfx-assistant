import io
import zipfile

import pytest
from docx import Document
from reportlab.pdfgen import canvas

from app import config
from app.services import document_ingestion as di
from tests.quote_fixtures import make_xlsx


def pdf_with_text(lines):
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    y = 800
    for line in lines:
        c.drawString(50, y, line)
        y -= 20
    c.save()
    return buf.getvalue()


def blank_pdf():
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    c.rect(50, 50, 200, 200, fill=1)
    c.save()
    return buf.getvalue()


def test_unsupported_extensions_get_a_helpful_message():
    with pytest.raises(di.IngestionError, match="Save it as .xlsx"):
        di.validate_upload("q.xls", b"data", set())
    with pytest.raises(di.IngestionError, match="Supported"):
        di.validate_upload("q.png", b"data", set())


def test_empty_oversized_fake_and_protected_files_are_rejected(monkeypatch):
    with pytest.raises(di.IngestionError, match="empty"):
        di.validate_upload("q.xlsx", b"", set())
    monkeypatch.setattr(config, "MAX_FILE_MB", 0)
    with pytest.raises(di.IngestionError, match="larger"):
        di.validate_upload("q.xlsx", b"x", set())
    monkeypatch.undo()
    with pytest.raises(di.IngestionError, match="real PDF"):
        di.validate_upload("q.pdf", b"not a pdf", set())
    with pytest.raises(di.IngestionError, match="real .xlsx"):
        di.validate_upload("q.xlsx", b"plain text", set())
    zipped = io.BytesIO()
    with zipfile.ZipFile(zipped, "w") as z:
        z.writestr("hello.txt", "x")
    with pytest.raises(di.IngestionError, match="real .docx"):
        di.validate_upload("q.docx", zipped.getvalue(), set())
    with pytest.raises(di.IngestionError, match="password"):
        di.validate_upload("q.xlsx", b"\xd0\xcf\x11\xe0" + b"0" * 20, set())


def test_duplicate_files_are_detected_by_content():
    data = make_xlsx([["a", 1, 2]])
    digest = di.validate_upload("one.xlsx", data, set())
    with pytest.raises(di.IngestionError, match="already"):
        di.validate_upload("renamed.xlsx", data, {digest})


def test_xlsx_rows_carry_references_and_hidden_content_is_ignored():
    data = make_xlsx([["Item", "Qty"], ["Pump", 2], [None, None], ["Valve", 10]], hidden_sheet=[["secret cost", 1]])
    doc = di.parse_document("q.xlsx", data)
    assert [r.ref for r in doc.rows] == ["Sheet1!R1", "Sheet1!R2", "Sheet1!R4"]
    assert doc.rows[1].text == "Pump | 2"
    assert doc.hidden_sheets == ["Internal"] and "secret" not in doc.as_text()


def test_xlsx_hidden_rows_are_skipped():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["a", 1])
    ws.append(["hidden row", 2])
    ws.append(["c", 3])
    ws.row_dimensions[2].hidden = True
    buf = io.BytesIO()
    wb.save(buf)
    assert "hidden row" not in di.parse_document("q.xlsx", buf.getvalue()).as_text()


def test_xlsx_formulas_without_a_saved_result_are_counted():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["a", 2, 3, "=B1*C1"])
    buf = io.BytesIO()
    wb.save(buf)
    assert di.parse_document("q.xlsx", buf.getvalue()).uncached_formulas == 1


def test_docx_paragraphs_and_tables_keep_their_order():
    d = Document()
    d.add_paragraph("Quotation letter")
    t = d.add_table(rows=2, cols=3)
    for r, vals in enumerate([["Item", "Qty", "Rate"], ["Pump", "2", "47,800"]]):
        for c, v in enumerate(vals):
            t.rows[r].cells[c].text = v
    d.add_paragraph("Freight extra.")
    buf = io.BytesIO()
    d.save(buf)
    doc = di.parse_document("q.docx", buf.getvalue())
    assert [r.ref for r in doc.rows] == ["P1", "T1-R1", "T1-R2", "P2"]
    assert doc.rows[2].text == "Pump | 2 | 47,800" and doc.tabular_rows == 1


def test_csv_and_tsv():
    doc = di.parse_document("q.csv", b"Item,Qty,Rate\nPump,2,47800\n\nValve,10,2240\n")
    assert [r.text for r in doc.rows] == ["Item | Qty | Rate", "Pump | 2 | 47800", "Valve | 10 | 2240"]
    assert di.parse_document("q.tsv", b"a\tb\tc\n1\t2\t3\n").rows[1].text == "1 | 2 | 3"


def test_pdf_text_is_read_and_image_only_pdf_is_detected():
    text_doc = di.parse_document("q.pdf", pdf_with_text(["Quotation No 12", "Pump 2 nos Rs 47,000 each. A longer line of text for the page."]))
    assert not text_doc.read_from_image and "Quotation No 12" in text_doc.as_text()
    assert text_doc.rows[0].ref.startswith("P1-B")
    assert di.parse_document("q.pdf", blank_pdf()).read_from_image


def test_corrupt_content_is_reported_not_raised():
    with pytest.raises(di.IngestionError, match="could not be read"):
        di.parse_document("q.pdf", b"%PDF-1.4 garbage")

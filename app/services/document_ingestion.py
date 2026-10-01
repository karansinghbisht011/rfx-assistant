"""Check uploaded quotation files and turn them into text rows with stable references.

References look like `Offer!R12` (sheet and row), `P4` (Word paragraph), `T1-R3` (Word table row),
`P1-B2` (PDF page 1, text block 2) or `P1-T1R5` (PDF table row). Gemini copies them back, and code
verifies them (verifier Q9). Files that cannot be used raise IngestionError with a reason a buyer can act on.
"""

import csv
import hashlib
import io
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

from app import config

ALLOWED_EXTENSIONS = (".csv", ".tsv", ".xlsx", ".pdf", ".docx")
SUPPORTED_TEXT = "CSV, TSV, Excel (.xlsx), PDF and Word (.docx)"
_OLE = b"\xd0\xcf\x11\xe0"


class IngestionError(Exception):
    """A file that cannot be processed; the message is shown to the buyer."""


@dataclass
class Row:
    ref: str
    text: str


@dataclass
class ParsedDocument:
    kind: str
    rows: list[Row] = field(default_factory=list)
    read_from_image: bool = False
    hidden_sheets: list[str] = field(default_factory=list)
    uncached_formulas: int = 0
    tabular_rows: int = 0          # rows that look like table lines, for the completeness check (Q13)
    page_count: int = 0

    def as_text(self) -> str:
        return "\n".join(f"[{r.ref}] {r.text}" for r in self.rows)

    def row_map(self) -> dict[str, str]:
        return {r.ref: r.text for r in self.rows}


def file_hash(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def validate_upload(filename: str, data: bytes, seen_hashes: set[str]) -> str:
    """Checks that need no parsing (Q1-Q3). Returns the content hash, or raises IngestionError."""
    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        hint = " Save it as .xlsx or .docx and upload again." if ext in (".xls", ".doc") else ""
        raise IngestionError(f"{ext or 'This'} files are not supported.{hint} Supported: {SUPPORTED_TEXT}.")
    if not data:
        raise IngestionError("The file is empty.")
    if len(data) > config.MAX_FILE_MB * 1024 * 1024:
        raise IngestionError(f"The file is larger than {config.MAX_FILE_MB} MB.")
    if data[:4] == _OLE and ext in (".xlsx", ".docx"):
        raise IngestionError("The file is password-protected or in an old format. Remove the protection and try again.")
    if ext == ".pdf" and not data.startswith(b"%PDF-"):
        raise IngestionError("This is not a real PDF file, although it has a .pdf name.")
    if ext in (".xlsx", ".docx"):
        member = "xl/workbook.xml" if ext == ".xlsx" else "word/document.xml"
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                if member not in z.namelist():
                    raise IngestionError(f"This is not a real {ext} file, although it has that name.")
        except zipfile.BadZipFile as exc:
            raise IngestionError(f"This is not a real {ext} file, although it has that name.") from exc
    digest = file_hash(data)
    if digest in seen_hashes:
        raise IngestionError("You have already added this exact file.")
    return digest


def _clean(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        return str(int(value)) if value == int(value) else f"{value:.6f}".rstrip("0").rstrip(".")
    if hasattr(value, "isoformat"):
        return value.isoformat()[:10]
    return re.sub(r"\s+", " ", str(value)).strip()


def _looks_tabular(cells: list[str]) -> bool:
    filled = [c for c in cells if c]
    return len(filled) >= 3 and any(re.search(r"\d", c) for c in filled)


def parse_document(filename: str, data: bytes) -> ParsedDocument:
    ext = Path(filename).suffix.lower()
    try:
        if ext in (".csv", ".tsv"):
            return _parse_delimited(data, "\t" if ext == ".tsv" else ",")
        if ext == ".xlsx":
            return _parse_xlsx(data)
        if ext == ".docx":
            return _parse_docx(data)
        return _parse_pdf(data)
    except IngestionError:
        raise
    except Exception as exc:  # a corrupt or odd file: report it, never crash the page
        raise IngestionError("The file could not be read. It may be corrupt or password-protected.") from exc


def _parse_delimited(data: bytes, delimiter: str) -> ParsedDocument:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("latin-1")
    doc = ParsedDocument(kind="delimited")
    for n, cells in enumerate(csv.reader(io.StringIO(text), delimiter=delimiter), start=1):
        cells = [_clean(c) for c in cells]
        if not any(cells):
            continue
        if n > config.MAX_SHEET_ROWS:
            raise IngestionError(f"The file has more than {config.MAX_SHEET_ROWS} rows.")
        doc.rows.append(Row(f"R{n}", " | ".join(cells)))
        doc.tabular_rows += _looks_tabular(cells)
    return doc


def _parse_xlsx(data: bytes) -> ParsedDocument:
    from openpyxl import load_workbook

    values = load_workbook(io.BytesIO(data), data_only=True)
    formulas = load_workbook(io.BytesIO(data), data_only=False)
    doc = ParsedDocument(kind="xlsx")
    for ws in values.worksheets:
        if ws.sheet_state != "visible":
            doc.hidden_sheets.append(ws.title)
            continue
        if ws.max_row and ws.max_row > config.MAX_SHEET_ROWS:
            raise IngestionError(f"The sheet '{ws.title}' has more than {config.MAX_SHEET_ROWS} rows.")
        fws = formulas[ws.title]
        for r in range(1, ws.max_row + 1):
            if ws.row_dimensions[r].hidden:
                continue
            cells = [_clean(c.value) for c in ws[r]]
            for cell in fws[r]:
                if isinstance(cell.value, str) and cell.value.startswith("=") and ws.cell(r, cell.column).value is None:
                    doc.uncached_formulas += 1
            if not any(cells):
                continue
            while cells and not cells[-1]:
                cells.pop()
            doc.rows.append(Row(f"{ws.title}!R{r}", " | ".join(cells)))
            doc.tabular_rows += _looks_tabular(cells)
    return doc


def _parse_docx(data: bytes) -> ParsedDocument:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(io.BytesIO(data))
    doc = ParsedDocument(kind="docx")
    p_no = t_no = 0
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            text = re.sub(r"\s+", " ", Paragraph(child, document).text).strip()
            if text:
                p_no += 1
                doc.rows.append(Row(f"P{p_no}", text))
        elif tag == "tbl":
            t_no += 1
            for r_no, row in enumerate(Table(child, document).rows, start=1):
                cells, seen = [], None
                for cell in row.cells:
                    text = _clean(cell.text)
                    if cell._tc is not seen:  # merged cells repeat the same text
                        cells.append(text)
                    seen = cell._tc
                if any(cells):
                    doc.rows.append(Row(f"T{t_no}-R{r_no}", " | ".join(cells)))
                    doc.tabular_rows += _looks_tabular(cells)
    return doc


def _overlap(a, b) -> float:
    """Share of rectangle a covered by rectangle b."""
    x0, y0, x1, y1 = max(a[0], b[0]), max(a[1], b[1]), min(a[2], b[2]), min(a[3], b[3])
    area = max(a[2] - a[0], 1e-6) * max(a[3] - a[1], 1e-6)
    return max(0.0, x1 - x0) * max(0.0, y1 - y0) / area


def _parse_pdf(data: bytes) -> ParsedDocument:
    import pymupdf

    pdf = pymupdf.open(stream=data, filetype="pdf")
    if pdf.needs_pass:
        raise IngestionError("The PDF is password-protected. Remove the protection and try again.")
    if pdf.page_count > config.MAX_PDF_PAGES:
        raise IngestionError(f"The PDF has more than {config.MAX_PDF_PAGES} pages.")
    doc = ParsedDocument(kind="pdf", page_count=pdf.page_count)
    chars = 0
    for page in pdf:
        n = page.number + 1
        items: list[tuple[float, str, str]] = []
        tables = []
        try:
            tables = page.find_tables().tables
        except Exception:
            tables = []
        for t_no, table in enumerate(tables, start=1):
            for r_no, cells in enumerate(table.extract(), start=1):
                cleaned = [_clean(c) for c in cells]
                if any(cleaned):
                    items.append((table.bbox[1] + r_no * 0.01, f"P{n}-T{t_no}R{r_no}", " | ".join(cleaned)))
                    doc.tabular_rows += _looks_tabular(cleaned)
        b_no = 0
        for block in page.get_text("blocks"):
            if block[6] != 0 or any(_overlap(block[:4], t.bbox) > 0.5 for t in tables):
                continue
            text = re.sub(r"\s+", " ", block[4]).strip()
            if text:
                b_no += 1
                items.append((block[1], f"P{n}-B{b_no}", text))
        for _, ref, text in sorted(items, key=lambda i: i[0]):
            doc.rows.append(Row(ref, text))
            chars += len(text)
    if chars < 30 * max(pdf.page_count, 1):
        doc.read_from_image = True  # little or no text layer: the file is sent to Gemini as the PDF itself
    return doc
